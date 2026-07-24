"""Shared Word monthly report layout and matplotlib charts (Cortex / Stellar Darktrace)."""

from __future__ import annotations

import logging
import os
import re
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import matplotlib

matplotlib.use("Agg")
import matplotlib.font_manager as fm
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import FancyBboxPatch

from app.report.logo import DEFAULT_HEADER_LOGO, DEFAULT_SOURCE_LOGO, ensure_watermark_png

logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)

REPO_ROOT = Path(__file__).resolve().parents[2]
WATERMARK_WIDTH_CM = 14.0
TOP_INCIDENTS_COUNT = 10

C_PRIMARY = "#29333A"
C_ACCENT = "#F37021"
C_LIGHT = "#A5B1C2"
C_RED = "#D83127"
C_AMBER = "#F5A623"
C_GREEN = "#4CAF50"
C_GRAY = "#7F7F7F"

CN_FONT = "Microsoft JhengHei"
EN_FONT = "Calibri"

RGB_PRIMARY = RGBColor(0x29, 0x33, 0x3A)
RGB_ACCENT = RGBColor(0xF3, 0x70, 0x21)
RGB_RED = RGBColor(0xD8, 0x31, 0x27)
RGB_AMBER = RGBColor(0xF5, 0xA6, 0x23)
RGB_GREEN = RGBColor(0x4C, 0xAF, 0x50)
RGB_GRAY = RGBColor(0x7F, 0x7F, 0x7F)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RGB_595959 = RGBColor(0x59, 0x59, 0x59)

HEX_PRIMARY = "29333A"
HEX_ACCENT = "F37021"
HEX_RED = "D83127"
HEX_AMBER = "F5A623"
HEX_GREEN = "4CAF50"
HEX_GRAY = "F2F2F2"

PRIORITY_COLOR = {"高": RGB_RED, "中": RGB_AMBER, "低": RGB_GREEN}
RISK_FILL = {"高": HEX_RED, "中": HEX_AMBER, "低": HEX_GREEN}
RISK_LABEL = {"高": "高 HIGH", "中": "中 MEDIUM", "低": "低 LOW"}


def setup_cjk_font() -> None:
    candidates = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
        "C:/Windows/Fonts/msjh.ttc",
        "C:/Windows/Fonts/mingliu.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ]
    for p in candidates:
        if os.path.exists(p):
            try:
                fm.fontManager.addfont(p)
            except Exception:
                pass
    plt.rcParams["font.family"] = [
        "Noto Sans CJK JP",
        "Noto Sans CJK TC",
        "Droid Sans Fallback",
        "Microsoft JhengHei",
        "PingFang TC",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def make_kpi_chart(kpis, out_path: Path) -> None:
    n = len(kpis)
    fig, axes = plt.subplots(1, n, figsize=(3.4 * n, 3.6))
    fig.patch.set_facecolor("white")
    if n == 1:
        axes = [axes]

    for i, ax in enumerate(axes):
        item = kpis[i]
        label, cur, prev, good_down = item[:4]
        subtitle = item[4] if len(item) > 4 else ""

        trend_text = "—"
        trend_color = C_GRAY
        if isinstance(cur, (int, float)) and isinstance(prev, (int, float)) and prev not in (0, "N/A"):
            delta = (float(cur) - float(prev)) / float(prev) * 100 if float(prev) else 0
            arrow = "↑" if delta >= 0 else "↓"
            trend_text = f"{arrow} {abs(delta):.1f}% MoM"
            is_good = (delta < 0 and good_down) or (delta > 0 and not good_down)
            trend_color = C_GREEN if is_good else C_RED
        elif subtitle and not isinstance(cur, (int, float)):
            trend_text = subtitle
            subtitle = ""

        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis("off")
        ax.add_patch(
            FancyBboxPatch(
                (0.02, 0.04),
                0.96,
                0.92,
                boxstyle="round,pad=0.02,rounding_size=0.05",
                linewidth=1.2,
                edgecolor="#BFBFBF",
                facecolor="#F8F9FA",
            )
        )
        ax.add_patch(
            FancyBboxPatch(
                (0.02, 0.90),
                0.96,
                0.06,
                boxstyle="round,pad=0.0,rounding_size=0.02",
                linewidth=0,
                facecolor=C_ACCENT,
            )
        )
        ax.text(0.5, 0.78, label, ha="center", va="center", fontsize=9, color=C_PRIMARY, weight="bold")

        main = f"{cur:,}" if isinstance(cur, int) else str(cur)
        ax.text(0.5, 0.48, main, ha="center", va="center", fontsize=22, color=C_PRIMARY, weight="bold")

        if subtitle:
            ax.text(0.5, 0.22, subtitle, ha="center", va="center", fontsize=7, color=C_GRAY)
        elif isinstance(prev, (int, float)) and prev != "N/A":
            prev_str = f"{int(prev):,}" if isinstance(prev, int) else str(prev)
            ax.text(0.5, 0.22, f"上期 Last: {prev_str}", ha="center", va="center", fontsize=8, color=C_GRAY)
        elif isinstance(prev, str) and prev not in ("N/A", ""):
            ax.text(0.5, 0.22, f"上期 Last: {prev}", ha="center", va="center", fontsize=8, color=C_GRAY)

        ax.text(0.5, 0.08, trend_text, ha="center", va="center", fontsize=8, color=trend_color, weight="bold")

    plt.suptitle("關鍵營運指標 Key Performance Indicators", fontsize=13, color=C_PRIMARY, weight="bold", y=1.02)
    plt.tight_layout(rect=[0, 0, 1, 0.94])
    plt.savefig(out_path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()


def make_risk_chart(risk_dist: dict[str, int], out_path: Path, *, basis: str = "Cases") -> None:
    fig, ax = plt.subplots(1, 1, figsize=(6, 4))
    fig.patch.set_facecolor("white")

    sizes = [risk_dist["高"], risk_dist["中"], risk_dist["低"]]
    labels = [f"高風險 High ({sizes[0]:,})", f"中風險 Medium ({sizes[1]:,})", f"低風險 Low ({sizes[2]:,})"]
    colors = [C_RED, C_AMBER, C_GREEN]
    total = sum(sizes)

    if total > 0:
        wedges, _ = ax.pie(
            sizes,
            colors=colors,
            startangle=90,
            counterclock=False,
            wedgeprops=dict(width=0.42, edgecolor="white", linewidth=2),
        )
        ax.legend(wedges, labels, loc="lower center", bbox_to_anchor=(0.5, -0.12), fontsize=9, frameon=False, ncol=1)
        ax.text(0, 0.08, f"{total:,}", ha="center", va="center", fontsize=22, color=C_PRIMARY, weight="bold")
        ax.text(0, -0.14, "Total", ha="center", va="center", fontsize=10, color=C_GRAY)
    else:
        ax.axis("off")
        ax.text(0.5, 0.5, "無資料", ha="center", va="center", fontsize=14, color=C_GRAY)

    ax.set_title(f"整體風險分佈 Risk Distribution ({basis})", fontsize=12, color=C_PRIMARY, weight="bold", pad=10)
    plt.tight_layout()
    plt.savefig(out_path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()


def _set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_para_border(paragraph, edge: str, color: str, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    p_bdr.append(el)


def _set_run_font(run, size_pt=11, bold=False, italic=False, color=None, cn_font=CN_FONT, en_font=EN_FONT):
    run.font.name = en_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), en_font)
    r_fonts.set(qn("w:hAnsi"), en_font)
    r_fonts.set(qn("w:eastAsia"), cn_font)


def add_run(paragraph, text, **kw):
    run = paragraph.add_run(text)
    _set_run_font(run, **kw)
    return run


def set_cell_text(cell, text, *, size_pt=10, bold=False, color=None, align="left", fill=None):
    cell.text = ""
    if fill:
        _set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    for j, line in enumerate(str(text).split("\n")):
        if j > 0:
            para.add_run().add_break()
        add_run(para, line, size_pt=size_pt, bold=bold, color=color)


def set_table_borders(table, color="BFBFBF", size="4"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tbl_borders.append(b)
    tbl_pr.append(tbl_borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def add_heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    add_run(p, text, size_pt=15, bold=True, color=RGB_PRIMARY)
    _set_para_border(p, "bottom", HEX_ACCENT, size="12")
    return p


def add_risk_cell(cell, level: str):
    set_cell_text(
        cell,
        RISK_LABEL[level],
        size_pt=10,
        bold=True,
        color=RGB_WHITE,
        align="center",
        fill=RISK_FILL[level],
    )


def _render_top_incident_cell(
    cell,
    header: str,
    value: str,
    *,
    fill: str | None = None,
    action_color: RGBColor | None = None,
) -> None:
    h = str(header or "").strip()
    if h == "風險":
        add_risk_cell(cell, value)
    elif h == "處理進度":
        set_cell_text(cell, value or " ", size_pt=9, align="center", fill=fill)
    elif h == "#":
        set_cell_text(cell, value, size_pt=10, bold=True, align="center", fill=fill)
    elif h in ("發生日期", "時間"):
        set_cell_text(cell, value, size_pt=9, align="center", fill=fill)
    elif h in ("攻擊類型", "模組", "類型", "等級"):
        set_cell_text(cell, value, size_pt=9, bold=True, color=RGB_PRIMARY, align="center", fill=fill)
    elif h in ("Case 狀態", "處置結果 Action", "處置結果"):
        set_cell_text(
            cell,
            value,
            size_pt=9,
            bold=True,
            color=action_color or RGB_GREEN,
            align="center",
            fill=fill,
        )
    else:
        set_cell_text(cell, value, size_pt=9, fill=fill)


def add_event_deep_dive_section(doc, data: dict) -> None:
    dives = data.get("event_deep_dives") or []
    if not dives:
        return

    add_heading(doc, data.get("event_deep_dive_heading", "五、事件報告說明"))
    field_rows = [
        ("案件名稱", "case_name", False),
        ("觸發時間", "trigger_time", False),
        ("主機名稱", "hostname", False),
        ("IP", "ip", False),
        ("事件說明", "event_description", True),
        ("原因分析", "root_cause_analysis", True),
    ]
    label_width = 3.2
    value_width = 13.8

    for idx, dive in enumerate(dives, start=1):
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(8 if idx > 1 else 2)
        title_p.paragraph_format.space_after = Pt(4)
        add_run(title_p, f"事件 {idx}", size_pt=11, bold=True, color=RGB_PRIMARY)

        tbl = doc.add_table(rows=len(field_rows), cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)
        for row_idx, (label, key, blank_area) in enumerate(field_rows):
            label_cell = tbl.rows[row_idx].cells[0]
            value_cell = tbl.rows[row_idx].cells[1]
            set_cell_text(
                label_cell,
                label,
                size_pt=10,
                bold=True,
                color=RGB_WHITE,
                fill=HEX_PRIMARY,
            )
            value = dive.get(key, "")
            if blank_area and not str(value).strip():
                value = "\n\n\n\n"
            set_cell_text(value_cell, value, size_pt=10)
        set_col_widths(tbl, [label_width, value_width])
        doc.add_paragraph()


def add_page_number_field(paragraph, instr: str):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(b)
    run._r.append(it)
    run._r.append(e)
    _set_run_font(run, size_pt=8, color=RGB_GRAY)


def flatten_image_for_word(src: Path, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_mtime >= src.stat().st_mtime:
        return dest
    try:
        from PIL import Image
    except ImportError:
        import shutil

        shutil.copy2(src, dest)
        return dest
    img = Image.open(src).convert("RGBA")
    bg = Image.new("RGB", img.size, (255, 255, 255))
    bg.paste(img, mask=img.split()[3])
    bg.save(dest, "JPEG", quality=92)
    return dest


def sanitize_docx_package(path: Path) -> None:
    tmp_fd, tmp_name = tempfile.mkstemp(suffix=".docx")
    os.close(tmp_fd)
    tmp_path = Path(tmp_name)
    try:
        with ZipFile(path, "r") as zin, ZipFile(tmp_path, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/stylesWithEffects.xml":
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b'<Override PartName="/word/stylesWithEffects.xml" '
                        b'ContentType="application/vnd.ms-word.stylesWithEffects+xml"/>',
                        b"",
                    )
                elif item.filename == "word/_rels/document.xml.rels":
                    data = re.sub(
                        br'<Relationship[^>]*Target="stylesWithEffects\.xml"[^>]*/>\s*',
                        b"",
                        data,
                    )
                zout.writestr(item, data)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def fix_drawing_docpr_ids(doc: Document) -> None:
    doc_elements = [doc.element.body]
    for section in doc.sections:
        doc_elements.append(section.header._element)
        doc_elements.append(section.first_page_header._element)
        doc_elements.append(section.footer._element)
    docprs: list = []
    for root in doc_elements:
        docprs.extend(root.xpath(".//wp:docPr"))
    for idx, doc_pr in enumerate(docprs, start=1):
        doc_pr.set("id", str(idx))


def add_anchor_watermark_to_header(header, watermark_path: Path, *, width_cm: float = 14.0) -> None:
    wp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    wp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    wp.paragraph_format.space_before = Pt(0)
    wp.paragraph_format.space_after = Pt(0)
    wp.paragraph_format.line_spacing = Pt(0.1)

    run = wp.add_run()
    run.add_picture(str(watermark_path), width=Cm(width_cm))

    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        return
    extent = inline.find(qn("wp:extent"))
    if extent is None:
        return
    cx, cy = extent.get("cx"), extent.get("cy")
    blip = drawing.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    if blip is None:
        return
    r_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    run._r.remove(drawing)

    anchor = parse_xml(
        f'<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" '
        f'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="9999" name="JJNET Watermark"/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="JJNET"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{r_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        f"</pic:pic></a:graphicData></a:graphic>"
        f"</wp:anchor></w:drawing>"
    )
    run._r.append(anchor)


def add_mitre_summary_table(doc, mitre_summary: dict) -> None:
    rows = mitre_summary.get("rows") or []
    cols = mitre_summary.get("columns") or ["主要 MITRE 戰術", "最高嚴重度", "原始告警", "聚合事件"]
    widths = mitre_summary.get("col_widths_cm") or [5.2, 2.2, 2.2, 2.2]

    tbl = doc.add_table(rows=1 + max(len(rows), 1), cols=len(cols))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    for i, h in enumerate(cols):
        set_cell_text(
            tbl.rows[0].cells[i],
            h,
            size_pt=10,
            bold=True,
            color=RGB_WHITE,
            align="center",
            fill=HEX_PRIMARY,
        )

    if rows:
        for r, row in enumerate(rows, start=1):
            fill = HEX_GRAY if r % 2 == 0 else None
            tactic, sev, issues_n, cases_n = row[:4]
            set_cell_text(tbl.rows[r].cells[0], tactic, size_pt=10, bold=True, color=RGB_PRIMARY, fill=fill)
            sev_color = RGB_RED if sev == "高" else (RGB_AMBER if sev == "中" else None)
            set_cell_text(
                tbl.rows[r].cells[1],
                sev,
                size_pt=10,
                bold=sev in ("高", "中"),
                color=sev_color,
                align="center",
                fill=fill,
            )
            set_cell_text(tbl.rows[r].cells[2], f"{issues_n:,}", size_pt=10, align="center", fill=fill)
            set_cell_text(tbl.rows[r].cells[3], f"{cases_n:,}", size_pt=10, align="center", fill=fill)
    else:
        empty_msg = mitre_summary.get("empty_message") or "（本期無含 MITRE 標籤之事件）"
        set_cell_text(tbl.rows[1].cells[0], empty_msg, size_pt=10, align="center")
        for i in range(1, len(cols)):
            set_cell_text(tbl.rows[1].cells[i], "—", size_pt=10, align="center")

    set_col_widths(tbl, widths)

    foot = mitre_summary.get("footnote") or ""
    if foot:
        fp = doc.add_paragraph()
        fp.paragraph_format.space_before = Pt(4)
        fp.paragraph_format.space_after = Pt(8)
        add_run(fp, foot, size_pt=9, italic=True, color=RGB_GRAY)


def cleanup_embed_artifacts(work_dir: Path) -> None:
    for path in sorted(work_dir.glob("_embed_*")):
        if not path.is_file():
            continue
        try:
            path.unlink()
        except OSError:
            pass


def add_cover_page(doc, *, cover: dict[str, str]) -> None:
    for _ in range(6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    t1 = doc.add_paragraph()
    t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t1.paragraph_format.space_before = Pt(48)
    add_run(t1, cover.get("title", ""), size_pt=22, bold=True, color=RGB_PRIMARY)

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(t2, cover.get("subtitle", "月度資安營運報告"), size_pt=16, bold=True, color=RGB_ACCENT)

    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(t3, cover.get("subtitle_en", "Monthly Security Operations Report"), size_pt=13, color=RGB_595959)

    tagline = cover.get("tagline", "")
    if tagline:
        t4 = doc.add_paragraph()
        t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t4.paragraph_format.space_before = Pt(8)
        add_run(t4, tagline, size_pt=10, italic=True, color=RGB_595959)

    doc.add_page_break()


def build_monthly_report(
    data: dict,
    kpi_chart: Path,
    risk_chart: Path | None,
    out_path: Path,
    *,
    logo_path: Path | None = None,
    repo_root: Path | None = None,
) -> None:
    root = repo_root or REPO_ROOT
    logo_src = Path(logo_path).expanduser() if logo_path else DEFAULT_SOURCE_LOGO
    if not logo_src.is_absolute():
        logo_src = root / logo_src
    watermark_img = ensure_watermark_png(
        source=logo_src if logo_src.is_file() else DEFAULT_HEADER_LOGO,
    )

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)

    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(11)

    confidential_line = data.get("confidential_line", "Confidential — MDR Monthly Report")

    def _add_confidential_header_line(header) -> None:
        hp = header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)
        add_run(hp, confidential_line, size_pt=8, italic=True, color=RGB_GRAY)

    def _setup_section_headers(section, watermark_path: Path, *, width_cm: float = 14.0) -> None:
        section.different_first_page_header_footer = True
        section.header.is_linked_to_previous = False
        section.first_page_header.is_linked_to_previous = False
        for hdr in (section.first_page_header, section.header):
            add_anchor_watermark_to_header(hdr, watermark_path, width_cm=width_cm)
            _add_confidential_header_line(hdr)

    def _add_section_footer(section) -> None:
        fp = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(fp, "Page ", size_pt=8, color=RGB_GRAY)
        add_page_number_field(fp, "PAGE")
        add_run(fp, " of ", size_pt=8, color=RGB_GRAY)
        add_page_number_field(fp, "NUMPAGES")

    for section in doc.sections:
        _add_section_footer(section)

    add_cover_page(doc, cover=data.get("cover") or {})

    info = doc.add_table(rows=3, cols=4)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info)
    rows_data = [
        ("客戶名稱", data["client_name"], "報告期間", data["period"]),
        ("服務窗口", data["service_owner"], "報告版本", data["version"]),
        ("資料來源", data["platform"], "整體風險評等", None),
    ]
    for r, row in enumerate(rows_data):
        cells = info.rows[r].cells
        for i, txt in enumerate(row):
            if i % 2 == 0:
                set_cell_text(cells[i], txt, size_pt=10, bold=True, color=RGB_WHITE, fill=HEX_PRIMARY)
            else:
                if r == 2 and i == 3:
                    add_risk_cell(cells[i], data["overall_risk"])
                else:
                    set_cell_text(cells[i], txt, size_pt=10)
    set_col_widths(info, [3.6, 4.0, 3.6, 4.0])

    doc.add_paragraph()

    add_heading(doc, data.get("summary_heading", "一、本月執行摘要 Executive Summary"))
    exec_p = doc.add_paragraph()
    runs = data.get("executive_summary_runs") or []
    for txt, bold, color in runs:
        add_run(exec_p, txt, size_pt=11, bold=bold, color=color)

    pic1 = doc.add_paragraph()
    pic1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    work_dir = out_path.parent
    kpi_embed = flatten_image_for_word(Path(kpi_chart), work_dir / "_embed_kpi.jpg")
    pic1.add_run().add_picture(str(kpi_embed), width=Cm(16.5))

    model_stats = data.get("model_alerts_stats")
    if model_stats:
        add_heading(doc, data.get("model_alerts_heading", "二、事件等級類別及事件分數統計"))
        foot = model_stats.get("footnote") or ""
        if foot:
            note_p = doc.add_paragraph()
            note_p.paragraph_format.space_after = Pt(6)
            add_run(note_p, foot, size_pt=9, italic=True, color=RGB_GRAY)
        panel_chart = data.get("model_alerts_panel_chart")
        if panel_chart:
            pic_ma = doc.add_paragraph()
            pic_ma.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_ma.paragraph_format.space_before = Pt(4)
            ma_embed = flatten_image_for_word(
                Path(panel_chart),
                work_dir / "_embed_model_alerts.jpg",
            )
            pic_ma.add_run().add_picture(str(ma_embed), width=Cm(16.5))
        doc.add_paragraph()

    add_heading(doc, data.get("mitre_heading", "二、MITRE 戰術摘要 MITRE Tactic Summary"))
    add_mitre_summary_table(doc, data.get("mitre_summary") or {"rows": [], "footnote": ""})

    if risk_chart is not None:
        pic2 = doc.add_paragraph()
        pic2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic2.paragraph_format.space_before = Pt(6)
        risk_embed = flatten_image_for_word(Path(risk_chart), work_dir / "_embed_risk.jpg")
        pic2.add_run().add_picture(str(risk_embed), width=Cm(10))

    top_n = data.get("top_count", TOP_INCIDENTS_COUNT)
    top_title = data.get("top_section_title", f"三、本月重點事件 Top {top_n} Incidents")
    add_heading(doc, top_title)
    cols = data.get("top_columns") or ["#", "發生日期", "模組", "事件描述 Description", "處置結果 Action", "風險"]
    widths = data.get("top_col_widths_cm") or [0.8, 2.0, 2.0, 6.6, 2.6, 1.4]
    incidents = data.get("incidents") or []
    tbl = doc.add_table(rows=1 + len(incidents), cols=len(cols))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    for i, h in enumerate(cols):
        set_cell_text(tbl.rows[0].cells[i], h, size_pt=10, bold=True, color=RGB_WHITE, align="center", fill=HEX_PRIMARY)
    for r, inc in enumerate(incidents, start=1):
        fill = HEX_GRAY if r % 2 == 0 else None
        action_color = data.get("top_action_color", RGB_GREEN)
        for c, header in enumerate(cols):
            val = inc[c] if c < len(inc) else ""
            _render_top_incident_cell(
                tbl.rows[r].cells[c],
                header,
                val,
                fill=fill,
                action_color=action_color,
            )
    set_col_widths(tbl, widths)

    doc.add_paragraph()

    add_event_deep_dive_section(doc, data)

    recommendations = data.get("recommendations") or []
    if recommendations:
        add_heading(doc, data.get("recommendations_heading", "四、改善建議 Recommendations"))
        for i, (title, body, prio) in enumerate(recommendations, start=1):
            p_t = doc.add_paragraph()
            p_t.paragraph_format.space_before = Pt(4)
            p_t.paragraph_format.space_after = Pt(2)
            label = f"{i}. {title}".rstrip(". ").rstrip()
            if not str(title or "").strip():
                label = f"{i}."
            add_run(p_t, label, size_pt=11, bold=True, color=RGB_PRIMARY)
            if str(prio or "").strip():
                add_run(p_t, "   [優先度: ", size_pt=9, color=RGB_GRAY)
                add_run(p_t, prio, size_pt=9, bold=True, color=PRIORITY_COLOR.get(prio, RGB_GRAY))
                add_run(p_t, "]", size_pt=9, color=RGB_GRAY)
            p_b = doc.add_paragraph()
            p_b.paragraph_format.space_after = Pt(6)
            add_run(p_b, body, size_pt=10)

    next_focus = data.get("next_focus") or []
    if next_focus:
        add_heading(doc, data.get("next_focus_heading", "五、下月服務重點 Next Month Focus"))
        for item in next_focus:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_run(p, item, size_pt=10)

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.paragraph_format.space_before = Pt(12)
    _set_para_border(end, "top", HEX_ACCENT, size="6")
    footer_note = data.get("footer_note", "— 本報告由 MDR SOC 自動產生，如有疑問請聯繫服務窗口 —")
    add_run(end, footer_note, size_pt=9, italic=True, color=RGB_GRAY)

    for section in doc.sections:
        _setup_section_headers(section, watermark_img, width_cm=WATERMARK_WIDTH_CM)
    fix_drawing_docpr_ids(doc)

    doc.save(out_path)
    sanitize_docx_package(out_path)
    cleanup_embed_artifacts(work_dir)
